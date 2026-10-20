"""Build the resQ 2.0 sales deck as BOTH a native .pptx and a matching .pdf.

Every relevant slide carries clickable links to the live persona pages and the
/demo console. Screenshots live in deck/shots/ (capture them first via the
screenshot step). Run:

    python -m scripts.build_deck

Outputs: deck/resQ_2.0_sales_deck.pptx  and  deck/resQ_2.0_sales_deck.pdf
"""
from __future__ import annotations

import base64
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(HERE, "deck", "shots")
OUT_DIR = os.path.join(HERE, "deck")

BASE = "https://resq-2-jrec.onrender.com"
REPO = "https://github.com/hansrajdeghun-ai/resq2.0"

# brand — "service ledger" light theme
INK = RGBColor(0x16, 0x23, 0x3F)      # ink navy
MARI = RGBColor(0xE8, 0x82, 0x0C)     # marigold accent
PAPER = RGBColor(0xF6, 0xF5, 0xF1)    # cool paper
CARD = RGBColor(0xFF, 0xFF, 0xFF)
MUTE = RGBColor(0x6B, 0x72, 0x84)
LINK = RGBColor(0xC2, 0x62, 0x00)

LINKS = [
    ("Live demo console", f"{BASE}/demo"),
    ("Customer app", f"{BASE}/customer"),
    ("Store owner", f"{BASE}/owner"),
    ("Field technician", f"{BASE}/field"),
    ("Area manager", f"{BASE}/area"),
    ("Admin (no-code)", f"{BASE}/admin"),
]

# ── slide content spec (shared by pptx + pdf) ───────────────────────────────
SLIDES = [
    {
        "kind": "cover",
        "eyebrow": "Reliance Digital · AI-native service",
        "title": "resQ 2.0",
        "subtitle": "The appliance-service agent that lives on WhatsApp.",
        "bullets": [
            "One AI agent · four WhatsApp numbers · thirteen journeys",
            "Event-sourced core · machine-first ledger · 8 protective rules",
        ],
        "links": [("▶  Open the live demo", f"{BASE}/demo"),
                  ("★  Source on GitHub", REPO)],
    },
    {
        "kind": "text",
        "eyebrow": "The problem",
        "title": "Service today leaks trust at every step",
        "bullets": [
            "Customers chase updates over calls — no single thread, no status.",
            "Jobs bounce between owner, technician and desk with no system of record.",
            "Cash paid before work is verified; disputes have no audit trail.",
            "Field edge-cases (wrong serial, no-show, damage) fall through the cracks.",
        ],
    },
    {
        "kind": "text",
        "eyebrow": "The solution",
        "title": "One agent on the channel people already use",
        "bullets": [
            "Customer, owner, field and area each get their own WhatsApp number.",
            "The agent books, assigns, tracks and closes — in plain language (incl. Hindi).",
            "Every action is an append-only event: the ledger is the source of truth.",
            "When a human must step in, the desk takes over the thread — then hands back.",
        ],
    },
    {
        "kind": "shot",
        "eyebrow": "Live demo",
        "title": "Four phones, one engine, in real time",
        "img": "demo.png",
        "note": "Install booked by the customer → offered to the owner → tracked by the area desk — all driven by the same event-sourced agent.",
        "links": [("Open /demo", f"{BASE}/demo")],
    },
    {
        "kind": "shot",
        "eyebrow": "Customer experience",
        "title": "A branded app that falls back to chat",
        "img": "customer.png",
        "note": "Repair & maintenance tiles open the WhatsApp thread and trigger the real journey intent. Photo, voice and Hindi all supported.",
        "links": [("Open /customer", f"{BASE}/customer")],
    },
    {
        "kind": "shot",
        "eyebrow": "Operations",
        "title": "Role consoles for owner, field & area",
        "img": "area.png",
        "note": "Area managers act only on exceptions; everything else the agent handles. Live jobs-by-state, action register (accept/reject), digest.",
        "links": [("/area", f"{BASE}/area"), ("/owner", f"{BASE}/owner"), ("/field", f"{BASE}/field")],
    },
    {
        "kind": "shot",
        "eyebrow": "Configuration",
        "title": "No-code control of the agent",
        "img": "admin.png",
        "note": "Onboard people, edit the rate card and routing/diagnosis as human-friendly tables — no deploy. Rules & SLAs stay hard-coded for safety.",
        "links": [("Open /admin", f"{BASE}/admin")],
    },
    {
        "kind": "grid",
        "eyebrow": "Coverage",
        "title": "All thirteen journeys, end to end",
        "items": [
            "J1 Install", "J2 Repair", "J3 Maintenance", "J4 AMC / plan",
            "J5 Reschedule", "J6 Spare part", "J7 Payment", "J8 Warranty",
            "J9 Relocation", "J10 Complaint", "J11 Feedback", "J12 Roster",
            "J13 Area ops",
        ],
        "note": "Plus shared assign/visit flows, the §4.1 exception table and an 11-timer scheduler.",
    },
    {
        "kind": "fullimg",
        "img": "architecture.png",
    },
    {
        "kind": "text",
        "eyebrow": "Trust by design",
        "title": "Eight protective rules the agent cannot break",
        "bullets": [
            "Technicians can never take cash in hand.",
            "No payout is released before the cooling-off window.",
            "Serial mismatch halts the job and alerts the desk.",
            "Every suggestion the agent makes is logged for accept/reject review.",
        ],
        "note": "Prices, SLAs, event types and the state machine are hard-coded — never AI-generated.",
    },
    {
        "kind": "links",
        "eyebrow": "Try it yourself",
        "title": "Every surface is live right now",
        "note": "Tap any link — the deployed app resets cleanly and runs the real engine.",
    },
]


# ── architecture diagram (dark hero, rendered to a PNG) ─────────────────────
ARCH_HTML = r"""
<html><head><meta charset='utf-8'><style>
* { box-sizing:border-box; margin:0; padding:0; font-family:'Helvetica Neue',Arial,sans-serif; }
body { width:1280px; height:720px; overflow:hidden; }
.canvas { position:relative; width:1280px; height:720px;
  background:
    radial-gradient(1100px 500px at 88% -8%, rgba(232,130,12,.30), transparent 60%),
    radial-gradient(900px 600px at 0% 110%, rgba(50,120,220,.18), transparent 55%),
    linear-gradient(160deg,#0e1730 0%,#16233f 55%,#0c1428 100%);
  color:#eaf0fb; padding:30px 52px 26px; }
.canvas::before { content:""; position:absolute; inset:0;
  background-image:radial-gradient(rgba(255,255,255,.055) 1px, transparent 1px);
  background-size:26px 26px; pointer-events:none; }
.eyebrow { color:#f6a743; font-weight:800; letter-spacing:.16em; font-size:13px; }
h1 { font-size:31px; margin-top:4px; letter-spacing:-.01em; }
h1 .accent { color:#f6a743; }
.sub { color:#9fb0cc; font-size:13px; margin-top:3px; }
.stack { position:relative; margin-top:11px; display:flex; flex-direction:column; gap:6px; z-index:2; }
.row { display:flex; align-items:stretch; gap:10px; }
.band { flex:1; display:flex; gap:10px; align-items:stretch;
  background:rgba(255,255,255,.045); border:1px solid rgba(255,255,255,.10);
  border-radius:16px; padding:8px 12px; backdrop-filter:blur(2px); }
.lbl { width:92px; flex:0 0 92px; display:flex; flex-direction:column; justify-content:center;
  color:#8aa0c4; font-size:10px; font-weight:800; letter-spacing:.14em; line-height:1.35; }
.cards { flex:1; display:flex; gap:8px; flex-wrap:wrap; align-items:center; }
.card { background:rgba(255,255,255,.07); border:1px solid rgba(255,255,255,.14);
  border-radius:11px; padding:7px 12px; font-size:12.5px; font-weight:650; color:#eef3fc; }
.card .k { display:block; font-size:10px; color:#9fb0cc; font-weight:700; letter-spacing:.04em; margin-top:2px; }
.card.hot { background:linear-gradient(180deg, rgba(232,130,12,.28), rgba(232,130,12,.14));
  border-color:rgba(246,167,67,.55); }
.card.core { background:rgba(80,150,250,.14); border-color:rgba(120,175,255,.4); }
.pill { background:rgba(255,255,255,.05); border:1px solid rgba(255,255,255,.12);
  border-radius:999px; padding:6px 12px; font-size:12px; color:#dbe6fb; font-weight:600; }
.arrow { position:relative; height:11px; display:flex; align-items:center; justify-content:center; }
.arrow span { color:#f6a743; font-size:11px; letter-spacing:9px; opacity:.8; }
.flowto { font-size:12px; color:#f6a743; font-weight:800; align-self:center; padding:0 2px; }
.guard { flex:0 0 250px; display:flex; flex-direction:column; gap:7px;
  background:linear-gradient(180deg, rgba(232,130,12,.16), rgba(232,130,12,.05));
  border:1px solid rgba(246,167,67,.45); border-radius:16px; padding:11px 13px; }
.guard .gt { font-size:10px; font-weight:800; letter-spacing:.14em; color:#f6a743; }
.guard .g { font-size:12.5px; color:#f3e7d4; font-weight:650; }
.corewrap { flex:1; display:flex; flex-direction:column; gap:7px; }
.coreflow { display:flex; align-items:center; gap:6px; }
.foot { position:absolute; left:52px; bottom:22px; color:#7e92b6; font-size:11px; z-index:2; }
.badge { position:absolute; right:52px; bottom:20px; color:#9fb0cc; font-size:11px; z-index:2; }
.badge b { color:#f6a743; }
</style></head><body>
<div class="canvas">
  <div class="eyebrow">UNDER THE HOOD</div>
  <h1>One agent, <span class="accent">one source of truth</span></h1>
  <div class="sub">Event-sourced core · machine-first ledger · deterministic guardrails · provider-agnostic AI</div>

  <div class="stack">
    <div class="row"><div class="band">
      <div class="lbl">SURFACES</div>
      <div class="cards">
        <div class="card hot">WhatsApp ×4<span class="k">customer · owner · field · area</span></div>
        <div class="card">Customer app</div>
        <div class="card">Role consoles</div>
        <div class="card">Admin</div>
        <div class="card">Demo console</div>
      </div>
    </div></div>
    <div class="arrow"><span>▼▼▼</span></div>

    <div class="row"><div class="band">
      <div class="lbl">INGRESS</div>
      <div class="cards">
        <div class="card">WhatsApp Cloud API webhook</div>
        <div class="card">REST <b>/api/*</b></div>
        <div class="card">Simulator</div>
        <div class="card">Channel-agnostic formatter <span class="k">§2.5</span></div>
      </div>
    </div></div>
    <div class="arrow"><span>▼▼▼</span></div>

    <div class="row"><div class="band">
      <div class="lbl">AGENT<br>BRAIN</div>
      <div class="cards">
        <div class="card core">Intent router</div><div class="flowto">→</div>
        <div class="card core">Dispatcher</div><div class="flowto">→</div>
        <div class="card hot">LLM orchestrator<span class="k">Claude ▸ OpenAI fallback</span></div>
        <div class="pill">translate</div><div class="pill">vision</div><div class="pill">voice</div>
        <div class="pill">pre-diagnosis</div><div class="pill">machine resolver</div>
      </div>
    </div></div>
    <div class="arrow"><span>▼▼▼</span></div>

    <div class="row"><div class="band">
      <div class="lbl">JOURNEYS</div>
      <div class="cards">
        <div class="card">J1–J13 handlers</div>
        <div class="card">Shared assign / visit</div>
        <div class="card">Exception table <span class="k">§4.1</span></div>
      </div>
    </div></div>
    <div class="arrow"><span>▼▼▼</span></div>

    <div class="row">
      <div class="band">
        <div class="lbl">EVENT-<br>SOURCED<br>CORE</div>
        <div class="corewrap">
          <div class="coreflow">
            <div class="card hot">Append-only event log</div><div class="flowto">→</div>
            <div class="card">Projections</div><div class="flowto">→</div>
            <div class="card">Job state machine</div><div class="flowto">→</div>
            <div class="card core">Machine-first ledger</div>
          </div>
          <div class="coreflow" style="margin-top:2px">
            <span class="card" style="background:none;border-style:dashed;color:#9fb0cc">no payout before the cooling-off window · the ledger is the record</span>
          </div>
        </div>
      </div>
      <div class="guard">
        <div class="gt">GUARDRAILS</div>
        <div class="g">⬦ 8 protective rules the agent can't break</div>
        <div class="g">⬦ 11-timer SLA scheduler</div>
      </div>
    </div>
    <div class="arrow"><span>▼▼▼</span></div>

    <div class="row"><div class="band">
      <div class="lbl">ADAPTERS<br>& DATA</div>
      <div class="cards">
        <div class="card">Kapture REST <span class="k">10 endpoints · live UAT reads verified</span></div>
        <div class="card">GCP Pub/Sub consumer</div>
        <div class="card">SQLAlchemy async <span class="k">SQLite dev · Postgres prod</span></div>
        <div class="pill">mock-first behind adapters</div>
      </div>
    </div></div>
  </div>

  <div class="badge">Rules · prices · SLAs · event types — <b>deterministic, never AI-generated</b></div>
</div>
</body></html>
"""


def render_arch():
    out = os.path.join(SHOTS, "architecture.png")
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_context(viewport={"width": 1280, "height": 720},
                           device_scale_factor=2).new_page()
        pg.set_content(ARCH_HTML, wait_until="networkidle")
        pg.wait_for_timeout(400)
        pg.locator(".canvas").screenshot(path=out)
        b.close()
    print("arch ->", out)
    return out


# ── PPTX renderer ───────────────────────────────────────────────────────────
def _emu_w():
    return Emu(12192000)  # 13.333in (16:9)


def _emu_h():
    return Emu(6858000)   # 7.5in


def _box(slide, l, t, w, h, fill=None, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.adjustments[0] = 0.04
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def _text(slide, l, t, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, space=6.0, font="Arial"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0)
    first = True
    for item in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.space_after = Pt(space); p.space_before = Pt(0)
        txt, url = (item if isinstance(item, tuple) else (item, None))
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold; r.font.name = font
        r.font.color.rgb = (LINK if url else color)
        if url:
            r.hyperlink.address = url
    return tb


def render_pptx():
    prs = Presentation(); prs.slide_width = _emu_w(); prs.slide_height = _emu_h()
    blank = prs.slide_layouts[6]
    IN = 914400
    for s in SLIDES:
        sl = prs.slides.add_slide(blank)
        if s["kind"] == "fullimg":
            sl.shapes.add_picture(os.path.join(SHOTS, s["img"]), 0, 0,
                                  prs.slide_width, prs.slide_height)
            continue
        bg = _box(sl, 0, 0, prs.slide_width, prs.slide_height, fill=PAPER)
        # accent rail
        _box(sl, 0, 0, Emu(int(0.18 * IN)), prs.slide_height, fill=MARI)
        mx = Emu(int(0.7 * IN))
        # eyebrow
        _text(sl, mx, Emu(int(0.55 * IN)), Emu(int(8 * IN)), Emu(int(0.4 * IN)),
              [s["eyebrow"].upper()], size=12, color=MARI, bold=True)
        # title
        tsize = 40 if s["kind"] == "cover" else 30
        _text(sl, mx, Emu(int(0.95 * IN)), Emu(int(8.2 * IN)), Emu(int(1.3 * IN)),
              [s["title"]], size=tsize, color=INK, bold=True)

        if s["kind"] == "cover":
            _text(sl, mx, Emu(int(2.2 * IN)), Emu(int(8 * IN)), Emu(int(0.7 * IN)),
                  [s["subtitle"]], size=20, color=INK)
            _text(sl, mx, Emu(int(3.0 * IN)), Emu(int(8 * IN)), Emu(int(1.4 * IN)),
                  ["•  " + b for b in s["bullets"]], size=16, color=MUTE, space=10)
            _text(sl, mx, Emu(int(4.9 * IN)), Emu(int(10 * IN)), Emu(int(0.9 * IN)),
                  s["links"], size=17, bold=True, space=8)
            _text(sl, mx, Emu(int(6.6 * IN)), Emu(int(10 * IN)), Emu(int(0.4 * IN)),
                  [BASE], size=12, color=MUTE)

        elif s["kind"] in ("text",):
            _text(sl, mx, Emu(int(2.1 * IN)), Emu(int(11.5 * IN)), Emu(int(4 * IN)),
                  ["•  " + b for b in s["bullets"]], size=19, color=INK, space=14)
            if s.get("note"):
                _text(sl, mx, Emu(int(6.2 * IN)), Emu(int(11.5 * IN)), Emu(int(0.8 * IN)),
                      [s["note"]], size=14, color=MUTE)

        elif s["kind"] == "grid":
            cols = 4; cw = 2.7; chh = 0.75; gx = 0.2; gy = 0.2
            for i, it in enumerate(s["items"]):
                r, c = divmod(i, cols)
                x = Emu(int((0.7 + c * (cw + gx)) * IN))
                y = Emu(int((2.2 + r * (chh + gy)) * IN))
                _box(sl, x, y, Emu(int(cw * IN)), Emu(int(chh * IN)), fill=CARD, line=RGBColor(0xE4, 0xE1, 0xD8))
                _text(sl, Emu(int((0.7 + c * (cw + gx) + 0.18) * IN)), y, Emu(int((cw - 0.3) * IN)),
                      Emu(int(chh * IN)), [it], size=15, color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            if s.get("note"):
                _text(sl, mx, Emu(int(6.4 * IN)), Emu(int(11.5 * IN)), Emu(int(0.7 * IN)),
                      [s["note"]], size=14, color=MUTE)

        elif s["kind"] == "shot":
            img = os.path.join(SHOTS, s["img"])
            # right card with screenshot
            from PIL import Image
            iw, ih = Image.open(img).size
            maxw, maxh = 6.6, 4.4
            ar = iw / ih
            w = maxw; h = w / ar
            if h > maxh:
                h = maxh; w = h * ar
            ix = Emu(int((13.333 - 0.7 - w) * IN)); iy = Emu(int(2.0 * IN))
            _box(sl, Emu(int(((13.333 - 0.7 - w) - 0.1) * IN)), Emu(int(1.9 * IN)),
                 Emu(int((w + 0.2) * IN)), Emu(int((h + 0.2) * IN)), fill=CARD, line=RGBColor(0xE4, 0xE1, 0xD8))
            sl.shapes.add_picture(img, ix, iy, Emu(int(w * IN)), Emu(int(h * IN)))
            # left note + links
            _text(sl, mx, Emu(int(2.3 * IN)), Emu(int(5 * IN)), Emu(int(3 * IN)),
                  [s["note"]], size=17, color=INK, space=10)
            _text(sl, mx, Emu(int(5.4 * IN)), Emu(int(5 * IN)), Emu(int(1.4 * IN)),
                  [("→  " + t, u) for t, u in s["links"]], size=16, bold=True, space=8)

        elif s["kind"] == "links":
            if s.get("note"):
                _text(sl, mx, Emu(int(2.0 * IN)), Emu(int(11 * IN)), Emu(int(0.6 * IN)),
                      [s["note"]], size=16, color=MUTE)
            cols = 2; cw = 5.7; chh = 0.95; gx = 0.3; gy = 0.25
            for i, (label, url) in enumerate(LINKS):
                r, c = divmod(i, cols)
                x = Emu(int((0.7 + c * (cw + gx)) * IN))
                y = Emu(int((2.8 + r * (chh + gy)) * IN))
                _box(sl, x, y, Emu(int(cw * IN)), Emu(int(chh * IN)), fill=CARD, line=RGBColor(0xE4, 0xE1, 0xD8))
                tx = Emu(int((0.7 + c * (cw + gx) + 0.25) * IN))
                _text(sl, tx, y, Emu(int((cw - 0.4) * IN)), Emu(int(chh * IN)),
                      [(label, url), ("   " + url, url)], size=15, bold=True, anchor=MSO_ANCHOR.MIDDLE, space=2)

    out = os.path.join(OUT_DIR, "resQ_2.0_sales_deck.pptx")
    prs.save(out)
    print("pptx ->", out)
    return out


# ── PDF renderer (styled HTML -> Playwright print) ──────────────────────────
def _b64(path):
    with open(path, "rb") as f:
        return "data:image/png;base64," + base64.b64encode(f.read()).decode()


def _slide_html(s):
    if s["kind"] == "fullimg":
        return (f'<section class="slide fullimg">'
                f'<img src="{_b64(os.path.join(SHOTS, s["img"]))}"></section>')
    eyebrow = f'<div class="eyebrow">{s["eyebrow"].upper()}</div>'
    title = f'<h1 class="{ "cover" if s["kind"]=="cover" else "" }">{s["title"]}</h1>'
    body = ""
    if s["kind"] == "cover":
        body += f'<p class="sub">{s["subtitle"]}</p><ul class="muted">'
        body += "".join(f"<li>{b}</li>" for b in s["bullets"]) + "</ul>"
        body += '<div class="cta">' + "".join(
            f'<a href="{u}">{t}</a>' for t, u in s["links"]) + "</div>"
        body += f'<div class="foot">{BASE}</div>'
    elif s["kind"] == "text":
        body += '<ul class="big">' + "".join(f"<li>{b}</li>" for b in s["bullets"]) + "</ul>"
        if s.get("note"):
            body += f'<div class="note">{s["note"]}</div>'
    elif s["kind"] == "grid":
        body += '<div class="grid">' + "".join(f'<div class="cell">{i}</div>' for i in s["items"]) + "</div>"
        if s.get("note"):
            body += f'<div class="note">{s["note"]}</div>'
    elif s["kind"] == "shot":
        links = "".join(f'<a href="{u}">→ {t}</a>' for t, u in s["links"])
        body += (f'<div class="shotwrap"><div class="shotleft"><p class="shotnote">{s["note"]}</p>'
                 f'<div class="shotlinks">{links}</div></div>'
                 f'<div class="shotimg"><img src="{_b64(os.path.join(SHOTS, s["img"]))}"></div></div>')
    elif s["kind"] == "links":
        if s.get("note"):
            body += f'<div class="note" style="margin-top:0">{s["note"]}</div>'
        body += '<div class="linkgrid">' + "".join(
            f'<a class="linkcard" href="{u}"><b>{t}</b><span>{u}</span></a>' for t, u in LINKS) + "</div>"
    return f'<section class="slide"><div class="rail"></div><div class="pad">{eyebrow}{title}{body}</div></section>'


def render_pdf():
    css = """
    * { box-sizing: border-box; margin:0; padding:0; font-family: 'Helvetica Neue', Arial, sans-serif; }
    @page { size: 1280px 720px; margin:0; }
    body { background:#f6f5f1; }
    .slide { position:relative; width:1280px; height:720px; background:#f6f5f1; overflow:hidden; page-break-after:always; }
    .slide.fullimg { padding:0; } .slide.fullimg img { width:1280px; height:720px; display:block; }
    .rail { position:absolute; left:0; top:0; width:14px; height:720px; background:#e8820c; }
    .pad { padding:48px 56px; }
    .eyebrow { color:#e8820c; font-weight:800; letter-spacing:.12em; font-size:14px; }
    h1 { color:#16233f; font-size:38px; margin-top:10px; line-height:1.1; }
    h1.cover { font-size:72px; letter-spacing:-.02em; }
    .sub { color:#16233f; font-size:24px; margin-top:18px; }
    ul { margin:22px 0 0 20px; color:#16233f; }
    ul.muted { color:#6b7284; font-size:18px; } ul.muted li { margin:8px 0; }
    ul.big li { font-size:22px; margin:16px 0; line-height:1.35; }
    .note { margin-top:26px; color:#6b7284; font-size:16px; max-width:1050px; }
    .cta { margin-top:34px; display:flex; gap:16px; }
    .cta a { background:#16233f; color:#fff; padding:14px 22px; border-radius:10px; font-weight:700; font-size:18px; text-decoration:none; }
    .cta a:last-child { background:#e8820c; }
    .foot { position:absolute; left:56px; bottom:40px; color:#6b7284; font-size:14px; }
    .grid { margin-top:26px; display:grid; grid-template-columns:repeat(4,1fr); gap:14px; max-width:1120px; }
    .cell { background:#fff; border:1px solid #e4e1d8; border-radius:10px; padding:18px; font-weight:700; color:#16233f; font-size:17px; }
    .shotwrap { margin-top:24px; display:flex; gap:34px; align-items:flex-start; }
    .shotleft { flex:0 0 420px; }
    .shotnote { font-size:20px; color:#16233f; line-height:1.4; }
    .shotlinks { margin-top:26px; display:flex; flex-direction:column; gap:12px; }
    .shotlinks a { color:#c26200; font-weight:700; font-size:18px; text-decoration:none; }
    .shotimg { flex:1; }
    .shotimg img { width:100%; border:1px solid #e4e1d8; border-radius:12px; box-shadow:0 10px 30px rgba(22,35,63,.12); }
    .linkgrid { margin-top:26px; display:grid; grid-template-columns:1fr 1fr; gap:16px; max-width:1120px; }
    .linkcard { background:#fff; border:1px solid #e4e1d8; border-radius:12px; padding:20px 24px; text-decoration:none; display:flex; flex-direction:column; gap:6px; }
    .linkcard b { color:#16233f; font-size:20px; } .linkcard span { color:#c26200; font-size:14px; }
    """
    html = ("<html><head><meta charset='utf-8'><style>" + css + "</style></head><body>"
            + "".join(_slide_html(s) for s in SLIDES) + "</body></html>")
    out = os.path.join(OUT_DIR, "resQ_2.0_sales_deck.pdf")
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_context(device_scale_factor=2).new_page()
        pg.set_content(html, wait_until="networkidle")
        pg.pdf(path=out, width="1280px", height="720px", print_background=True,
               margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
        b.close()
    print("pdf  ->", out)
    return out


if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    render_arch()
    render_pptx()
    render_pdf()
